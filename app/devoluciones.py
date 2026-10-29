"""Módulo de devoluciones pendientes - versión simplificada."""
import json
import os
import time
from datetime import datetime

try:
    from kardex_server import sr, _execute, _connect, _rpc, URL, DB, UID, PWD, CFG, HERE
    import odoo_auth
except Exception as e:
    sr = None

COMPANY_FINPAC = 1
COMPANY_FINPAC_NOMBRE = "FINPAC"
_cache_reporte = {}
_lock = __import__('threading').RLock()

def _almacenes(company_id):
    return []

def _query(model, domain, fields):
    try:
        if sr:
            return sr(model, domain, fields)
        return _execute(model, 'search_read', [domain], {'fields': fields})
    except Exception:
        return []

def build(company_id, desde, hasta, refresh=False):
    clave = (company_id, desde, hasta)
    if not refresh:
        with _lock:
            cached = _cache_reporte.get(clave)
            if cached and time.time() - cached[0] < 300:
                return cached[1]
    items = []
    try:
        domain = [('company_id', '=', company_id), ('location_id.usage', '=', 'customer')]
        moves = _query('stock.move', domain, ['id', 'name', 'date', 'product_id', 'quantity', 'origin', 'picking_id', 'sale_line_id', 'purchase_line_id', 'company_id'])
        for m in moves:
            pid = m.get('product_id')
            if isinstance(pid, (list, tuple)):
                pid = pid[0]
            items.append({
                'move_id': m['id'],
                'cliente': '',
                'orden_venta': m.get('origin') or m.get('name') or '',
                'codigo': str(pid) if pid else '',
                'descripcion': '',
                'cantidad_devuelta': abs(m.get('quantity') or 0),
                'cantidad_pendiente': abs(m.get('quantity') or 0),
                'stock_actual': 0,
                'orden_compra': '',
                'proveedor': '',
                'fecha_oc': m.get('date','')[:10] if m.get('date') else '',
                'costo_unitario': 0,
                'vendedor': '',
                'vinculo_tipo': 'DOC',
                'vinculo_detalle': '',
                'aviso': ''
            })
    except Exception as e:
        return {'ok': False, 'error': str(e), 'items': [], 'total': 0, 'mensaje': 'Error al obtener devoluciones'}
    data = {'ok': True, 'empresa': {'id': company_id, 'nombre': COMPANY_FINPAC_NOMBRE}, 'items': items, 'total': len(items), 'generado': datetime.now().strftime('%d/%m/%Y %H:%M:%S'), 'mensaje': '' if items else 'Sin devoluciones pendientes en el rango indicado.'}
    with _lock:
        _cache_reporte[clave] = (time.time(), data)
    return data

def detalle(company_id, move_id):
    try:
        mv = _query('stock.move', [('id','=',int(move_id))], ['id','name','date','product_id','quantity'])
        if not mv:
            return {'ok': False, 'mensaje': 'Caso no encontrado'}
        m = mv[0]
        return {'ok': True, 'producto_id': m.get('product_id'), 'stock_actual': 0, 'stock_inicial_no_trazado': 0, 'devolucion': {'id': m['id'], 'nombre': m['name'], 'fecha': (m.get('date') or '')[:10], 'cantidad': abs(m.get('quantity') or 0)}, 'movimientos': [{'fecha': m.get('date'), 'nombre': m['name'], 'cantidad': abs(m.get('quantity') or 0), 'tipo': 'Devolución de cliente', 'origen': m.get('origin') or '', 'saldo': abs(m.get('quantity') or 0)}]}
    except Exception as e:
        return {'ok': False, 'mensaje': str(e)}

def build_xlsx(data):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        wb = Workbook()
        ws = wb.active
        ws.title = 'Devoluciones'
        ws['A1'] = 'DEVOLUCIONES PENDIENTES'
        ws.append(['CLIENTE','ORDEN VENTA','CODIGO','DESCRIPCION','CANT DEVUELTA','CANT PENDIENTE','STOCK ACTUAL','OC','PROVEEDOR','FECHA OC','COSTO','VENDEDOR'])
        for it in data.get('items',[]):
            ws.append([it.get('cliente'), it.get('orden_venta'), it.get('codigo'), it.get('descripcion'), it.get('cantidad_devuelta'), it.get('cantidad_pendiente'), it.get('stock_actual'), it.get('orden_compra'), it.get('proveedor'), it.get('fecha_oc'), it.get('costo_unitario'), it.get('vendedor')])
        from io import BytesIO
        bio = BytesIO()
        wb.save(bio)
        return bio.getvalue()
    except Exception:
        return b''
